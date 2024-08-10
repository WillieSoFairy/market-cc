import sys

sys.path.insert(0, "./pypkg")
import pymysql as mysql
import dbConfig as db
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Font, Border, Side
from datetime import datetime
import base64


def main(event, context):
    order_date = event["order_date"]
    batch = event["batch"]
    deliver_date = event["deliver_date"]
    try:
        order, export_info = query_db(order_date, batch)
    except IOError as err:
        return {"status": -1, "info": err}

    fileDate = datetime.strptime(order_date, r"%Y-%m-%d").strftime(r"%Y%m%d")
    fileNames = []
    try:
        for i, info in enumerate(export_info):
            orders = order[i]
            file_name = f"{orders[0][0]}_{orders[0][1]}_{fileDate}_{batch}.xlsx"
            wb = create_excel(orders, deliver_date, info)
            wb.save(f"/tmp/{file_name}")
            wb.close()
            fileNames.append(file_name)
    except Exception as err:
        return {"status": -1, "info": f"{err}-{i}"}

    files = []
    for i, name in enumerate(fileNames):
        with open(f"/tmp/{name}", "rb") as f:
            content = base64.b64encode(f.read()).decode("utf-8")
            files.append({"id": (export_info[i])[0], "name": name, "file": content})

    return {"files": files, "status": 0, "info": None, "order_date": fileDate}


def query_db(order_date, batch):
    try:
        conn = mysql.connect(
            host=db.host, user=db.user, passwd=db.passwd, database=db.database
        )
        with conn.cursor() as cursor:
            sql = "SELECT id,id_in_batch,deliver_ent,receiver,sender FROM export_rec WHERE order_date=%s AND batch=%s;"
            cursor.execute(sql, (order_date, batch))
            export_info = cursor.fetchall()

            data = []
            for x in export_info:
                sql = "select ent_name,dept_name,good_name,count,unit_name,remark from order_view2 where export_id=%s;"
                cursor.execute(sql, x[0])
                rows = cursor.fetchall()
                data.append(rows)
        return (data, export_info)
    except:
        raise IOError("Query database error")


def create_excel(orders: tuple[tuple], deliver_date, export_info):
    try:
        wb = Workbook()
        ent_name = (orders[0])[0]
        dept_name = (orders[0])[1]
        client = ent_name if dept_name == None else f"{ent_name}（{dept_name}）"
        ws = wb.create_sheet(client)
        write_header(ws, export_info[1], deliver_date, client, export_info[2])
        write_data(ws, orders)
        write_footer(ws, export_info[3], export_info[4])
        wb.remove(wb["Sheet"])
        return wb
    except:
        raise Exception("Create sheets error")


def write_header(ws: Worksheet, note_id, date, client, deliver_ent):
    ws.merge_cells(f"A1:G1")
    title = ws.cell(1, 1, "送货单")
    ws.append(
        [
            "单号：",
            str(note_id).zfill(2),
            "开单日期：",
            None,
            date,
        ]
    )
    ws.append(["客户：", client, "供货单位：", None, deliver_ent])
    for i in [3, 5]:
        offset = 1 if i != 5 else 2
        for j in range(1, 3):
            ws.merge_cells(
                start_column=i,
                end_column=i + offset,
                start_row=1 + j,
                end_row=1 + j,
            )

    title.alignment = Alignment(horizontal="center", vertical="center")
    title.font = Font(bold=True, size=14, name="黑体")
    for row in ws[f"A1:G1"]:
        for cell in row:
            cell.border = Border(bottom=Side("double"))

    for i in range(2, 4):
        for j in [1, 3]:
            label = ws.cell(i, j)
            label.font = Font(name="黑体")
            label.alignment = Alignment("right", "center")
        for j in [2, 5]:
            text = ws.cell(i, j)
            text.alignment = Alignment("left", "center")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 6


def write_data(ws: Worksheet, orders: tuple[tuple]):
    init = ws.max_row + 1
    ws.append(["编号", "菜品名称", "数量", "单位", "单价", "金额", "备注"])
    for index, row in enumerate(orders):
        ws.append(
            [
                index + 1,
                row[2],
                row[3],
                row[4],
                None,
                None,
                row[5],
            ]
        )
    for i in range(init, ws.max_row + 1):
        for j in range(1, 8):
            data = ws.cell(i, j)
            sideThin = Side("thin")
            data.border = Border(sideThin, sideThin, sideThin, sideThin)
            data.alignment = Alignment(vertical="center")
            if i == init:
                data.font = Font(name="黑体")
                data.alignment = Alignment("center", "bottom")


def write_footer(ws: Worksheet, receiver, sender):
    init = ws.max_row + 1
    ws.merge_cells(f"A{init}:E{init}")
    sum = ws.cell(init, 1, "小计")
    sum.alignment = Alignment("right", "bottom")
    sum.font = Font(name="黑体")
    sideThin = Side("thin")
    for row in ws[f"A{init}:G{init}"]:
        for cell in row:
            cell.border = Border(sideThin, sideThin, sideThin, sideThin)
    ws.append([None, "收货人：", receiver, None, "送货人：", sender])
    for i in [2, 5]:
        label = ws.cell(init + 1, i)
        label.alignment = Alignment("right", "bottom")
        label.font = Font(name="黑体")
        ws.merge_cells(
            start_column=i + 1, end_column=i + 2, start_row=init + 1, end_row=init + 1
        )
        text = ws.cell(init + 1, i + 1)
        text.alignment = Alignment("left", "bottom")
